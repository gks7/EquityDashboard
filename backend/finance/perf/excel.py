"""Excel export of the performance base (same content as the page, plus the daily detail)."""
import datetime as dt
import io

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .engine import MESES

NAVY = '1F3A5F'
HDR = PatternFill('solid', fgColor=NAVY)
SUB = PatternFill('solid', fgColor='E8EEF5')
EST = PatternFill('solid', fgColor='FFF7E0')
F = 'Arial'
PCT = '0.00%;[Red]-0.00%;-'
USD = '#,##0;[Red](#,##0);-'
USD2 = '#,##0.00;[Red](#,##0.00);-'
COTA = '0.000000'


def _v(v):
    if v is None:
        return None
    if isinstance(v, (float, np.floating)):
        return None if (np.isnan(v) or np.isinf(v)) else float(v)
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, pd.Timestamp):
        return v.date()
    if isinstance(v, str) and len(v) == 10 and v[4] == '-' and v[7] == '-':
        try:
            return dt.date.fromisoformat(v)
        except ValueError:
            return v
    return v


def _header(ws, row, n):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=F, bold=True, color='FFFFFF', size=10)
        cell.fill = HDR
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30


def _title(ws, text, sub=''):
    ws['A1'] = text
    ws['A1'].font = Font(name=F, bold=True, size=14, color=NAVY)
    if sub:
        ws['A2'] = sub
        ws['A2'].font = Font(name=F, italic=True, size=9, color='5A6B7D')


def _table(ws, df, r0, fmts=None, widths=None):
    fmts = fmts or {}
    cols = list(df.columns)
    for j, c in enumerate(cols, 1):
        ws.cell(row=r0, column=j, value=str(c))
    _header(ws, r0, len(cols))
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            v = _v(v)
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = Font(name=F, size=9)
            if cols[j - 1] in fmts:
                cell.number_format = fmts[cols[j - 1]]
            elif isinstance(v, dt.date):
                cell.number_format = 'dd/mm/yyyy'
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(c, max(10, min(40, len(str(c)) + 4)))
    ws.freeze_panes = ws.cell(row=r0 + 1, column=2)
    return r0 + len(df)


def build_workbook(payload, frames):
    wb = Workbook()
    k = payload['kpi']
    f = frames['f']

    ws = wb.active
    ws.title = 'Resumo'
    _title(ws, 'IGF WM Total Return Fund — Class A (série líder)', f"Gerado em {dt.date.today():%d/%m/%Y}. Valores em USD.")
    rows = [('Cota em ' + _v(k['date']).strftime('%d/%m/%Y'), k['cota'], COTA),
            ('Cota oficial ' + _v(payload['official_last']).strftime('%d/%m/%Y') + ' (administrador)', payload['official_cota'], COTA),
            ('Rentabilidade no mês', k['mtd'], PCT), ('Rentabilidade no ano', k['ytd'], PCT), ('Rentabilidade 12 meses', k['m12'], PCT),
            ('Rentabilidade desde o início', k['itd'], PCT), ('Rentabilidade anualizada', k['ann'], PCT),
            ('Volatilidade anualizada', k['vol'], PCT), ('Máximo drawdown', k['maxdd'], PCT),
            ('Patrimônio líquido (USD)', k['nav'], USD), ('Cotas (todas as séries)', k['shares'], '#,##0'),
            ("Marca d'água (performance)", k['hwm'], COTA)]
    ws['A4'], ws['B4'] = 'Indicador', 'Valor'
    _header(ws, 4, 2)
    for i, (lab, val, fm) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=lab).font = Font(name=F, size=10)
        c = ws.cell(row=i, column=2, value=val)
        c.number_format = fm
        c.font = Font(name=F, size=10, bold=True)
    ws.column_dimensions['A'].width = 46
    ws.column_dimensions['B'].width = 18
    r = 5 + len(rows) + 1
    ws.cell(row=r, column=1, value='Verificações').font = Font(name=F, bold=True, size=10, color=NAVY)
    for i, ch in enumerate(payload['checks'], r + 1):
        ws.cell(row=i, column=1, value=f"[{ch['status'].upper()}] {ch['title']}").font = Font(name=F, size=9)
        ws.cell(row=i, column=2, value=ch['detail']).font = Font(name=F, size=9, color='5A6B7D')

    # monthly returns
    ws = wb.create_sheet('Rentab_Mensal')
    _title(ws, 'Rentabilidade mensal', 'Meses fechados = cota oficial do administrador; mês corrente = estimativa.')
    months = pd.DataFrame(payload['months'])
    series_names = [c for c in months.columns if c != 'm']
    hdr = ['Série / Ano'] + MESES + ['Ano', 'Acumulado']
    for j, h in enumerate(hdr, 1):
        ws.cell(row=4, column=j, value=h)
    _header(ws, 4, len(hdr))
    rr = 5
    for name in series_names:
        ws.cell(row=rr, column=1, value=name).font = Font(name=F, bold=True, size=10, color=NAVY)
        rr += 1
        acc = 1.0
        for y in sorted(set(m[:4] for m in months.m)):
            g = months[months.m.str.startswith(y)]
            ws.cell(row=rr, column=1, value=y).font = Font(name=F, size=9, bold=True)
            yr = 1.0
            for _, row in g.iterrows():
                v = row[name]
                if v is not None and not pd.isna(v):
                    c = ws.cell(row=rr, column=1 + int(row.m[5:7]), value=float(v))
                    c.number_format = PCT
                    c.font = Font(name=F, size=9)
                    yr *= 1 + v
            acc *= yr
            for col, val in ((14, yr - 1), (15, acc - 1)):
                c = ws.cell(row=rr, column=col, value=val)
                c.number_format = PCT
                c.font = Font(name=F, size=9, bold=col == 14)
                c.fill = SUB
            rr += 1
        rr += 1
    ws.column_dimensions['A'].width = 24

    # daily cota
    ws = wb.create_sheet('Cota_Diaria')
    _title(ws, 'Cota diária, patrimônio e cotas', 'Retorno diário por fórmula. Linhas amarelas = estimativa.')
    cd = pd.DataFrame({'Data': [d.date() for d in f.index], 'Cota': f.cota.values, 'Retorno dia': None, 'PL (USD)': f.nav.values,
                       'Cotas': f.shares.values, 'Aplicações (USD)': f.subscription.values,
                       'Fonte': np.where(f.estimated, 'Estimada', np.where(f.official_daily, 'Oficial diária', 'Estimada (ajustada ao fechamento oficial)'))})
    end = _table(ws, cd, 4, {'Cota': COTA, 'Retorno dia': PCT, 'PL (USD)': USD, 'Cotas': '#,##0', 'Aplicações (USD)': USD}, {'Fonte': 40})
    for i in range(6, end + 1):
        ws.cell(row=i, column=3, value=f'=IF(B{i-1}=0,"",B{i}/B{i-1}-1)').number_format = PCT
        if ws.cell(row=i, column=7).value == 'Estimada':
            for j in range(1, 8):
                ws.cell(row=i, column=j).fill = EST

    # assets
    ws = wb.create_sheet('Ativos_Resumo')
    _title(ws, 'Performance por ativo', 'TWR; P&L inclui proventos; contribuição em pontos da cota.')
    a = pd.DataFrame(payload['assets'])
    cols = [('id', 'Ativo'), ('name', 'Nome'), ('bucket', 'Classe'), ('sector', 'Setor'), ('status', 'Status'), ('first', 'Entrada'),
            ('last', 'Última data'), ('units', 'Quantidade'), ('price', 'Preço'), ('mv', 'Valor (USD)'), ('w', '% PL'), ('cost', 'Custo médio'),
            ('gain', 'Ganho s/ custo'), ('r_mtd', 'Ret. mês'), ('r_3m', 'Ret. 3m'), ('r_ytd', 'Ret. ano'), ('spx_ytd', 'S&P mesma janela (ano)'),
            ('r_itd', 'Ret. desde entrada'), ('spx_itd', 'S&P desde entrada'), ('pnl_mtd', 'P&L mês'), ('pnl_ytd', 'P&L ano'), ('pnl_itd', 'P&L total'),
            ('inc', 'Proventos'), ('c_ytd', 'Contrib. ano'), ('c_itd', 'Contrib. total')]
    t = a[[c for c, _ in cols]].copy()
    t.columns = [n for _, n in cols]
    pf = {n: PCT for n in ['% PL', 'Ganho s/ custo', 'Ret. mês', 'Ret. 3m', 'Ret. ano', 'S&P mesma janela (ano)', 'Ret. desde entrada', 'S&P desde entrada', 'Contrib. ano', 'Contrib. total']}
    pf.update({n: USD for n in ['Valor (USD)', 'P&L mês', 'P&L ano', 'P&L total', 'Proventos']})
    pf.update({'Quantidade': '#,##0', 'Preço': '#,##0.00', 'Custo médio': '#,##0.00'})
    end = _table(ws, t, 4, pf, {'Nome': 34, 'Ativo': 22, 'Classe': 18, 'Setor': 22})
    names = list(t.columns)
    tr = end + 1
    ws.cell(row=tr, column=1, value='TOTAL').font = Font(name=F, bold=True, size=9)
    for col in ['Valor (USD)', '% PL', 'P&L mês', 'P&L ano', 'P&L total', 'Proventos', 'Contrib. ano', 'Contrib. total']:
        j = names.index(col) + 1
        L = get_column_letter(j)
        c = ws.cell(row=tr, column=j, value=f'=SUM({L}5:{L}{end})')
        c.number_format = pf[col]
        c.font = Font(name=F, bold=True, size=9)
        c.fill = SUB
    for col in ['Ret. ano', 'Ret. desde entrada', 'Ganho s/ custo']:
        L = get_column_letter(names.index(col) + 1)
        ws.conditional_formatting.add(f'{L}5:{L}{end}', ColorScaleRule(start_type='num', start_value=-0.3, start_color='F4B6B6',
                                                                       mid_type='num', mid_value=0, mid_color='FFFFFF',
                                                                       end_type='num', end_value=0.3, end_color='A9D8B8'))

    for sheet, key, fm in (('Ativos_Ret_Mensal', 'mret', PCT), ('Ativos_PnL_Mensal', 'mpnl', USD)):
        ws = wb.create_sheet(sheet)
        _title(ws, sheet.replace('_', ' '))
        m = frames[key].copy()
        order = [x['id'] for x in payload['assets'] if x['id'] in m.index]
        m = m.loc[order]
        m.columns = [f'{MESES[p.month-1]}/{str(p.year)[2:]}' for p in m.columns]
        m = m.reset_index().rename(columns={'asset_id': 'Ativo'})
        end = _table(ws, m, 4, {c: fm for c in m.columns[1:]}, {'Ativo': 24})
        if fm == USD:
            ws.cell(row=end + 1, column=1, value='TOTAL').font = Font(name=F, bold=True, size=9)
            for j in range(2, len(m.columns) + 1):
                L = get_column_letter(j)
                c = ws.cell(row=end + 1, column=j, value=f'=SUM({L}5:{L}{end})')
                c.number_format = USD
                c.font = Font(name=F, bold=True, size=9)

    ws = wb.create_sheet('Atribuicao')
    _title(ws, 'Atribuição mensal por classe (p.p. da cota)')
    c = frames['cls'].copy()
    c.index = [f'{MESES[p.month-1]}/{p.year}' for p in c.index]
    c = c.reset_index().rename(columns={'index': 'Mês'})
    _table(ws, c, 4, {k2: PCT for k2 in c.columns[1:]}, {'Mês': 12})

    ws = wb.create_sheet('Posicoes_Diarias')
    _title(ws, 'Posições e P&L diário por ativo')
    x = frames['pab'][['date', 'asset_id', 'bucket', 'units', 'price', 'mv_open', 'mv_close', 'net_invested', 'income', 'pnl', 'ret', 'contrib', 'source']].copy()
    x['date'] = x.date.dt.date
    x = x[(x.mv_close.abs() > 0) | (x.pnl.abs() > 0.01)].sort_values(['date', 'asset_id'])
    x.columns = ['Data', 'Ativo', 'Classe', 'Quantidade', 'Preço', 'Valor abertura', 'Valor fechamento', 'Investido no dia', 'Proventos', 'P&L', 'Retorno', 'Contribuição', 'Fonte']
    _table(ws, x, 4, {'Quantidade': '#,##0.##', 'Preço': '#,##0.0000', 'Valor abertura': USD, 'Valor fechamento': USD, 'Investido no dia': USD,
                      'Proventos': USD2, 'P&L': USD, 'Retorno': PCT, 'Contribuição': '0.000%'}, {'Ativo': 24, 'Fonte': 26, 'Classe': 18})
    ws.auto_filter.ref = f'A4:M{4 + len(x)}'

    ws = wb.create_sheet('Transacoes')
    _title(ws, 'Extrato classificado + ajustes manuais')
    led = frames['led'].copy()
    cols = [c for c in ['trade_date', 'booking_date', 'type', 'asset_id', 'units', 'price', 'accrued', 'amount', 'counterparty', 'is_reversal', 'source', 'desc1', 'txn_no'] if c in led.columns]
    led = led[cols].sort_values('trade_date')
    end = _table(ws, led, 4, {'units': '#,##0.##', 'amount': USD2, 'price': '#,##0.0000', 'accrued': USD2}, {'desc1': 60, 'counterparty': 34, 'asset_id': 24})
    ws.auto_filter.ref = f'A4:{get_column_letter(len(cols))}{end}'

    ws = wb.create_sheet('Verificacoes')
    _title(ws, 'Conciliação e verificações')
    ch = pd.DataFrame([(c['status'], c['title'], c['detail']) for c in payload['checks']], columns=['Status', 'Verificação', 'Detalhe'])
    _table(ws, ch, 4, widths={'Verificação': 50, 'Detalhe': 110})

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
