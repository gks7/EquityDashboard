import datetime as dt
import io

import openpyxl
import pandas as pd
from django.test import TestCase

from finance.perf.statement import AliasMatcher, classify, dealing_date, parse_workbook as parse_statement
from finance.perf.admin_report import parse_workbook as parse_admin, lead_month_ends
from finance.perf.engine import anchored, hwm_from

ALIASES = [
    dict(asset_id='QQQ', isin='US46090E1038', statement_pattern=r'Invesco QQQ', tickers='QQQ,QQQ US'),
    dict(asset_id='SDHA', isin='IE00BZ17CN18', statement_pattern='', tickers='SDHA LN EQUITY,SDHA LN'),
    dict(asset_id='UGPABZ 5.25 10/06/26', isin='USL9412AAA53', statement_pattern='Ultrapar', tickers=''),
]


def statement_xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in [('Account number:', '0308 00152481.60'), ('From:', dt.datetime(2026, 9, 1)), ('Until:', dt.datetime(2026, 10, 8)),
              ('Opening balance:', 0), ('Closing balance:', sum((r[5] or 0) + (r[6] or 0) for r in rows)), ('Valued in:', 'USD'), (None,)]:
        ws.append(list(r))
    ws.append(['Trade date', 'Trade time', 'Booking date', 'Value date', 'Currency', 'Debit', 'Credit', 'Individual amount',
               'Balance', 'Transaction no.', 'Description1', 'Description2', 'Description3', 'Footnotes'])
    for r in rows:
        ws.append(list(r))
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


class StatementParserTests(TestCase):
    def test_classification_units_and_mapping(self):
        d = dt.datetime(2026, 9, 1)
        rows = [
            (d, None, d, d, 'USD', -645331.53, None, None, 0, 'A1', 'Invesco QQQ Trust Series I; US46090E1038', 'Purchase Spot',
             'Number/Amt. 900; Transaction price: 716.963333 USD (*a); Transaction no. A1', None),
            (d, None, d, d, 'USD', 645331.53, None, None, 0, 'A2', 'Invesco QQQ Trust Series I; US46090E1038', 'Purchase Spot;Reversal',
             'Number/Amt. 900; Transaction price: 716.963333 USD; Transaction no. A2', None),
            (d, None, d, d, 'USD', None, 200000, None, 0, 'A3', '5.25%  Notes Ultrapar International S.A. 2016-06.10.2026 Reg S; USL9412AAA53',
             'Secur. Repay.', 'Number/Amt. 200000 USD; Transaction price: 100 %; Transaction no. A3', None),
            (d, None, d, d, 'USD', None, 3000000, None, 0, 'A4', 'BANCO BTG PACTUAL SA CAYMAN BRANCH;FLOOR 5', 'credit', 'Reason for payment: X', None),
            (d, None, d, d, 'USD', -34380.75, None, None, 0, 'A5', 'IGF WEALTH MANAGEMENT LTDA;AV BRIG FARIA LIMA', 'MGMT FEE AUG/2026; e-banking payment order', '', None),
            (d, None, d, d, 'USD', None, 400115.32, None, 0, 'A6', 'You sold CAD; You bought USD; FX CS-SGRFS', 'Sale FX Spot', 'You bought: 400115.32 USD', None),
            (d, None, d, d, 'USD', -8000, None, None, 0, 'A7', 'PRICEWATERHOUSECOOPERS AUDITORES IN;BR', 'PWC - AUDIT INV 139570; e-banking payment order', '', None),
        ]
        meta, out = parse_statement(statement_xlsx(rows), AliasMatcher(ALIASES))
        types = [r['type'] for r in out]
        self.assertEqual(types, ['BUY', 'BUY', 'REDEMPTION', 'SUBSCRIPTION', 'MGMT_FEE', 'TRANSFER', 'EXPENSE'])
        self.assertEqual(out[0]['units'], 900)
        self.assertEqual(out[1]['units'], -900)          # reversal cancels the booking
        self.assertTrue(out[1]['is_reversal'])
        self.assertEqual(out[0]['asset_id'], 'QQQ')
        self.assertEqual(out[2]['asset_id'], 'UGPABZ 5.25 10/06/26')
        self.assertEqual(out[2]['units'], -200000)
        self.assertEqual(meta['account'], '0308 00152481.60')

    def test_wrong_file_is_rejected(self):
        wb = openpyxl.Workbook()
        wb.active.append(['foo'])
        bio = io.BytesIO(); wb.save(bio); bio.seek(0)
        with self.assertRaises(ValueError):
            parse_statement(bio, AliasMatcher(ALIASES))

    def test_matcher_tickers(self):
        m = AliasMatcher(ALIASES)
        self.assertEqual(m.from_ticker('SDHA LN EQUITY'), 'SDHA')
        self.assertEqual(m.from_ticker('QQQ US'), 'QQQ')
        self.assertIsNone(m.from_ticker('ZZZ'))

    def test_dealing_date(self):
        self.assertEqual(dealing_date(dt.date(2026, 9, 28)), dt.date(2026, 9, 30))
        self.assertEqual(dealing_date(dt.date(2026, 4, 2)), dt.date(2026, 3, 31))
        self.assertEqual(dealing_date(dt.date(2026, 3, 12)), dt.date(2026, 2, 27))   # Feb-28 is a Saturday
        self.assertEqual(dealing_date(dt.date(2026, 8, 28)), dt.date(2026, 8, 31))


class AdminReportTests(TestCase):
    def _wb(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'NAV Report'
        for r in [['IGF WM TOTAL RETURN FUND LTD. SAC - Class A Segregated Account'], [dt.datetime(2026, 9, 30)], ['TOTAL FUND POSITION'],
                  ['Base NAV', 39923794.87, 42451734.22], ['  Management fee', -34380.75, -35378.21], ['  Performance fee', -694.71, 0],
                  ['NAV', 39888719.41, 42416356.01], ['Redemptions', 0, 0], ['Subscriptions', 3000000, 4000000], ['NAV Closing', 42888719.41, 46416356.01],
                  ['EQUITIES', 'CUSTODIAN', 'ISIN', '', 'CURRENCY', 'AMOUNT', 'COST VALUE', 'VALUE', 'PRICE', 'RATE', 'WEIGHT'],
                  [' POWERSHARES QQQ TRUST SERIES ', 'UBS', 'US46090E1038', None, 'USD', 7810, 4636928.73, 5777603.7, 739.77, 1, 0.124],
                  [5777603.7, 0.124],
                  ['BONDS', 'CUSTODIAN', 'ISIN', 'INTEREST', 'CURRENCY', 'AMOUNT', 'COST VALUE', 'VALUE', 'PRICE', 'RATE', 'WEIGHT'],
                  ['ULTRAPAR INTERNATIONAL SA', 'UBS', 'USL9412AAA53', 5104.17, 'USD', 200000, 198025, 199938, 97.4169, 1, 0.0043],
                  ['BANK ACCOUNTS', 'CUSTODIAN', '', '', 'CURRENCY', 'AMOUNT', '', 'VALUE', 'PRICE', 'RATE', 'WEIGHT'],
                  ['308-152481-01', 'UBS', 'USD', 3723416.87, 3723416.87, 1, 1, 0.08],
                  ['TOTAL', 46475107.1, 1.0]]:
            ws.append(r)
        ps = wb.create_sheet('Participating shares')
        for r in [['PARTICIPATING SHARES NOVEMBER 2025 SERIES'], ['NAV (CLOSING)', 32208632.5, 1.12635, 36278296.2],
                  ['Year', 'Jan', 'Feb'], ['2026', 1.1183, 1.1109], ['%', 0.0078, -0.0066], ['2025', 1.1088, 1.1096],
                  ['PARTICIPATING SHARES SEPTEMBER 2026 SERIES'], ['SUBSCRIPTIONS', 3588829.4, 1.11441, 4000000], ['NAV (CLOSING)', 3588829.4, 1.11441, 4000000]]:
            ps.append(r)
        bio = io.BytesIO(); wb.save(bio); bio.seek(0)
        return bio

    def test_parse(self):
        p = parse_admin(self._wb(), 'x.xlsx')
        self.assertEqual(p['date'], dt.date(2026, 9, 30))
        self.assertAlmostEqual(p['NAV Closing'], 46416356.01)
        self.assertAlmostEqual(p['Subscriptions_prev'], 3000000)
        self.assertAlmostEqual(p['Management fee'], -35378.21)
        secs = [h['section'] for h in p['holdings']]
        self.assertEqual(secs, ['EQUITIES', 'BONDS', 'BANK ACCOUNTS'])
        self.assertEqual(p['holdings'][1]['accrued'], 5104.17)
        self.assertAlmostEqual(p['total_shares'], 32208632.5 + 3588829.4)
        me = lead_month_ends([p])
        self.assertAlmostEqual(me[dt.date(2026, 1, 30)], 1.1183)
        self.assertAlmostEqual(me[dt.date(2025, 11, 28)], 1.1088)   # series launched in November
        self.assertAlmostEqual(me[dt.date(2026, 9, 30)], 1.12635)


class CotaMathTests(TestCase):
    def test_anchored_hits_month_ends(self):
        days = pd.bdate_range('2026-01-01', '2026-03-10').date
        r = pd.Series(0.001, index=days)
        anchors = pd.Series({dt.date(2025, 12, 31): 1.0, dt.date(2026, 1, 30): 1.01, dt.date(2026, 2, 27): 1.03})
        cota, adj = anchored(r, anchors)
        self.assertAlmostEqual(cota[dt.date(2026, 1, 30)], 1.01)
        self.assertAlmostEqual(cota[dt.date(2026, 2, 27)], 1.03)
        self.assertGreater(cota[dt.date(2026, 3, 10)], 1.03)        # after the last anchor: pure estimate
        self.assertEqual(len(adj), 2)

    def test_hwm(self):
        anchors = {dt.date(2025, 11, 28): 1.1088, dt.date(2026, 3, 31): 1.20, dt.date(2026, 5, 29): 1.1414}
        self.assertAlmostEqual(hwm_from(anchors, 1.1364), 1.1414)
